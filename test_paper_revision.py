from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from labot.paper_revision import create_word_table


def _revision_cells(tmp_path: Path, comment: str, response: str = ""):
    output = tmp_path / "revision.docx"
    create_word_table([("1", comment, response, "open")], output)
    cells = Document(output).tables[0].rows[1].cells
    return cells[1], cells[2]


def _number_start(paragraph) -> int:
    num_id = paragraph._p.pPr.numPr.numId.val
    numbering = paragraph.part.numbering_part.element
    num = next(item for item in numbering.num_lst if item.numId == num_id)
    return int(num.lvlOverride_lst[0].startOverride.val)


def test_numbered_list_after_prose_creates_distinct_paragraphs(tmp_path):
    markdown = """Strengths:
1. The paper addresses a highly practical problem.
2. The motivation is clearly articulated.
3. The artefact is novel.
4. The method is rigorous.
5. The findings are clear.
6. The writing is concise."""

    comment, _ = _revision_cells(tmp_path, markdown)

    assert [paragraph.text for paragraph in comment.paragraphs] == [
        "Strengths:",
        "The paper addresses a highly practical problem.",
        "The motivation is clearly articulated.",
        "The artefact is novel.",
        "The method is rigorous.",
        "The findings are clear.",
        "The writing is concise.",
    ]
    assert all(p.style.name == "List Number" for p in comment.paragraphs[1:])
    assert [_number_start(p) for p in comment.paragraphs[1:]] == list(range(1, 7))


def test_unordered_list_and_inline_formatting_work_in_response(tmp_path):
    _, response = _revision_cells(
        tmp_path,
        "Comment",
        "Actions:\n- Use **clear language**.\n* Add *more evidence*.\n+ Proofread.",
    )

    assert [p.text for p in response.paragraphs] == [
        "Actions:",
        "Use clear language.",
        "Add more evidence.",
        "Proofread.",
    ]
    assert all(p.style.name == "List Bullet" for p in response.paragraphs[1:])
    assert response.paragraphs[1].runs[1].bold
    assert response.paragraphs[2].runs[1].italic


def test_wrapped_prose_paragraphs_and_hard_breaks_are_preserved(tmp_path):
    comment, _ = _revision_cells(
        tmp_path,
        "A source-wrapped sentence\ncontinues in one paragraph.\n\n"
        "A hard break follows.  \nStill the same paragraph.",
    )

    assert len(comment.paragraphs) == 2
    assert comment.paragraphs[0].text == (
        "A source-wrapped sentence continues in one paragraph."
    )
    assert (
        comment.paragraphs[1].text == "A hard break follows.\nStill the same paragraph."
    )
    assert comment.paragraphs[1]._p.xpath(".//w:br")


def test_independent_ordered_lists_restart_at_source_number(tmp_path):
    comment, response = _revision_cells(
        tmp_path,
        "1) First comment\n2) Second comment",
        "1. First response\n2. Second response",
    )

    for cell in (comment, response):
        assert [p.style.name for p in cell.paragraphs] == ["List Number"] * 2
        assert [_number_start(p) for p in cell.paragraphs] == [1, 2]
        num_ids = [p._p.pPr.numPr.numId.val for p in cell.paragraphs]
        assert len(set(num_ids)) == 2


def test_numbering_is_saved_as_numbering_properties_not_literal_text(tmp_path):
    comment, _ = _revision_cells(tmp_path, "3. An item")

    paragraph = comment.paragraphs[0]
    assert paragraph.text == "An item"
    assert paragraph._p.pPr.numPr is not None
    assert paragraph._p.pPr.numPr.numId.get(qn("w:val")) is not None
    assert _number_start(paragraph) == 3
